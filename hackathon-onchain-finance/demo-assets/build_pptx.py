#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成一份可编辑的 AvaVault Perp 黑客松 Pitch Deck（.pptx），
文字是真正的 PowerPoint 文本框（可以直接改字），截图是真实运行截图（可以替换）。
与 build_slides.py 生成的视频用的是同一份文案/数据，但这里是「可编辑」版本。

用法: python3 build_pptx.py [zh|en]
"""
import sys
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

LANG = sys.argv[1] if len(sys.argv) > 1 else "zh"

BG = RGBColor(0x0B, 0x0D, 0x10)
CARD = RGBColor(0x15, 0x18, 0x1C)
RED = RGBColor(0xE8, 0x41, 0x42)
TEXT = RGBColor(0xEE, 0xEE, 0xEE)
MUTED = RGBColor(0x9A, 0xA4, 0xAF)
GREEN = RGBColor(0x2E, 0xCC, 0x71)

SCREEN_DIR = "./raw_screenshots"

CONTENT = {
  "zh": {
    "deck_title": "AvaVault Perp — RWA 抵押的永续合约保证金系统",
    "slides": [
      {"kicker": "AVALANCHE BUILDATHON · 链上金融赛道", "title": "AvaVault Perp",
       "bullets": ["把\u201c收租金的房子\u201d变成\u201c能交易永续合约的保证金\u201d",
                   "赛道：链上金融与交易（On-chain Finance & Trading）",
                   "状态：可运行的合约原型 + 完整测试 + 真实部署到 Fuji 测试网"]},
      {"kicker": "问题", "title": "真实世界资产，大多数时候是\u201c死资产\u201d",
       "bullets": ["RWA Token 持有者只能\u201c拿着等分红\u201d，流动性差、用途单一",
                   "永续合约交易者的保证金必须是稳定币/主流资产",
                   "收益类资产要先变现才能用作保证金——变现意味着卖出，失去持有收益权的敞口"]},
      {"kicker": "洞察 → 方案", "title": "只要有一个 DEX 交易对，就有一个可编程的价格",
       "bullets": ["存入 RWA Token → DEX 实时估值 → 打折计入保证金 → 开永续合约仓位",
                   "复用同一套 DEX Oracle 架构，为 RWA 资产提供实时估值",
                   "不需要先卖出 RWA，就能获得交易保证金"]},
      {"kicker": "技术融合", "title": "三层链上金融组件，拼成一套完整系统",
       "bullets": ["预言机层 · DEX Oracle — LFJ V1 实时报价，getAmountsOut 链上定价",
                   "资产层 · RWA Token — 房地产租金收益权，AccessControl 权限体系",
                   "结算层 · Perp Vault — 链下撮合 + 链上净额结算的永续合约引擎"]},
      {"kicker": "真实部署证明 · 非截图摆拍", "title": "MultiCollateralVault 已真实部署到 Fuji 测试网",
       "image": "02-vault-contract-snowtrace.png",
       "caption": "Snowtrace 可直接验证：合约创建交易 + 3 笔真实链上交互"},
      {"kicker": "真实链上交易", "title": "存入保证金 → 签名授权结算，全部可验证",
       "image2": ["03-real-deposit-tx-snowtrace.png", "04-real-settle-tx-snowtrace.png"],
       "caption2": ["100 枚 RWA Token 真实存入金库作为保证金", "EIP-712 签名授权，operator 代为完成真实结算"]},
      {"kicker": "实时仪表盘", "title": "没有一个数字是写死的",
       "image": "05-onchain-dashboard.png",
       "caption": "每次刷新都重新读取 Fuji 链上状态：供应量、托管余额、DEX 储备量估值"},
      {"kicker": "端到端 Demo", "title": "永续合约交易引擎，真实跑通",
       "image": "01-perp-dex-orderbook.png",
       "caption": "真实登录签名 · 真实订单簿深度 · 真实下单回报"},
      {"kicker": "为什么是 Avalanche", "title": "亚秒级确认 + 极低 Gas，让高频链上定价成为可能",
       "bullets": ["真实交易手续费：约 0.0001 AVAX 量级",
                   "DEX 实时价格可以被频繁、低成本地读取",
                   "C-Chain 完全 EVM 兼容，Solidity 直接复用"]},
      {"kicker": "AvaVault Perp", "title": "让沉睡的真实世界资产，重新流动起来",
       "bullets": ["GitHub: <填入你 fork 后的仓库链接>",
                   "赛道：链上金融（On-chain Finance）",
                   "感谢观看 · 我们在 Avalanche 上，构建下一个链上金融的可能"]},
    ],
  },
  "en": {
    "deck_title": "AvaVault Perp — RWA-Collateralized Perpetual Margin System",
    "slides": [
      {"kicker": "AVALANCHE BUILDATHON · ONCHAIN FINANCE TRACK", "title": "AvaVault Perp",
       "bullets": ["Turn a rent-collecting property into perpetual-trading margin",
                   "Track: On-chain Finance & Trading",
                   "Status: Working contract prototype + full tests + really deployed on Fuji testnet"]},
      {"kicker": "THE PROBLEM", "title": "Real-world assets are mostly \u201cdead assets\u201d on-chain",
       "bullets": ["RWA token holders can mostly only \u201chold and collect\u201d — poor liquidity, single use",
                   "Perpetual traders need stablecoin / blue-chip margin",
                   "Yield assets must be cashed out first to be used as margin — selling away the upside exposure"]},
      {"kicker": "INSIGHT -> SOLUTION", "title": "Any DEX pair gives you a programmable price",
       "bullets": ["Deposit RWA Token -> DEX real-time pricing -> Haircut into margin -> Open perp position",
                   "Reuses the same DEX oracle architecture to value RWA assets live",
                   "No need to sell the RWA token to get trading margin"]},
      {"kicker": "TECH FUSION", "title": "Three on-chain finance layers, fused into one system",
       "bullets": ["ORACLE · DEX Oracle — LFJ V1 live quotes via getAmountsOut on-chain pricing",
                   "ASSET · RWA Token — Real-estate rental rights, AccessControl role system",
                   "SETTLEMENT · Perp Vault — Off-chain matching + on-chain net settlement engine"]},
      {"kicker": "REAL DEPLOYMENT · NOT STAGED", "title": "MultiCollateralVault is really deployed on Fuji testnet",
       "image": "02-vault-contract-snowtrace.png",
       "caption": "Verifiable on Snowtrace: contract creation + 3 real on-chain interactions"},
      {"kicker": "REAL ON-CHAIN TRANSACTIONS", "title": "Deposit margin -> signed settlement, fully verifiable",
       "image2": ["03-real-deposit-tx-snowtrace.png", "04-real-settle-tx-snowtrace.png"],
       "caption2": ["100 RWA tokens really deposited as margin", "EIP-712 signed authorization, settled by the operator"]},
      {"kicker": "LIVE DASHBOARD", "title": "Nothing here is hardcoded",
       "image": "05-onchain-dashboard.png",
       "caption": "Every refresh re-reads Fuji chain state: supply, vault balance, DEX-reserve valuation"},
      {"kicker": "END-TO-END DEMO", "title": "The perpetual trading engine, really running",
       "image": "01-perp-dex-orderbook.png",
       "caption": "Real signed login - real order book depth - real order fills"},
      {"kicker": "WHY AVALANCHE", "title": "Sub-second finality + near-zero gas make this practical",
       "bullets": ["Real transaction fees: around 0.0001 AVAX",
                   "DEX prices can be read frequently, cheaply",
                   "C-Chain is fully EVM compatible, Solidity reused as-is"]},
      {"kicker": "AvaVault Perp", "title": "Putting sleeping real-world assets back to work",
       "bullets": ["GitHub: <put your forked repo link here>",
                   "Track: On-chain Finance",
                   "Thanks for watching - building the next wave of on-chain finance on Avalanche"]},
    ],
  },
}[LANG]


def set_background(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG


def add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return tb


def add_bullets(slide, left, top, width, height, items, size=20, color=TEXT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(14)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, s in enumerate(CONTENT["slides"], start=1):
        slide = prs.slides.add_slide(blank)
        set_background(slide)
        # top red bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Pt(6))
        bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()

        add_textbox(slide, Inches(0.7), Inches(0.35), Inches(11.9), Inches(0.4), s["kicker"], 15, RED, bold=True)
        add_textbox(slide, Inches(0.68), Inches(0.75), Inches(11.9), Inches(1.0), s["title"], 32, TEXT, bold=True)

        if "bullets" in s:
            add_bullets(slide, Inches(0.8), Inches(1.9), Inches(11.5), Inches(5), s["bullets"], size=20)
        elif "image" in s:
            pic_path = f"{SCREEN_DIR}/{s['image']}"
            pic = slide.shapes.add_picture(pic_path, Inches(2.0), Inches(1.9), height=Inches(4.6))
            add_textbox(slide, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.5), s["caption"], 14, MUTED, align=PP_ALIGN.CENTER)
        elif "image2" in s:
            slide.shapes.add_picture(f"{SCREEN_DIR}/{s['image2'][0]}", Inches(0.6), Inches(1.9), width=Inches(5.9))
            slide.shapes.add_picture(f"{SCREEN_DIR}/{s['image2'][1]}", Inches(6.8), Inches(1.9), width=Inches(5.9))
            add_textbox(slide, Inches(0.6), Inches(6.7), Inches(5.9), Inches(0.5), s["caption2"][0], 13, MUTED, align=PP_ALIGN.CENTER)
            add_textbox(slide, Inches(6.8), Inches(6.7), Inches(5.9), Inches(0.5), s["caption2"][1], 13, MUTED, align=PP_ALIGN.CENTER)

        add_textbox(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.35),
                    "AvaVault Perp · Avalanche Buildathon · Onchain Finance Track", 10, MUTED)
        add_textbox(slide, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.35), f"{idx}/10", 10, MUTED, align=PP_ALIGN.RIGHT)

    out = f"avavault-perp-pitch-{LANG}.pptx"
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    build()
